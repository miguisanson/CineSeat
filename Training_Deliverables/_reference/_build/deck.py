#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Render one .pptx per module from the shared slide spec.

Follows the reference deck's grammar: uppercase kicker -> title -> content, on a
near-black ground with a gold accent, everything absolutely positioned on Blank
layouts (layout index 6) for full control.
"""

import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import design as D

W, H = D.SLIDE_W_IN, D.SLIDE_H_IN
M = D.MARGIN_IN


def rgb(t):
    return RGBColor(*t)


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    f = s.background.fill
    f.solid()
    f.fore_color.rgb = rgb(D.BG)
    return s


def rect(s, x, y, w, h, colour):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(colour)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def rich(s, x, y, w, h, body, size, colour, lead_colour=None):
    """Renders '**Lead-in.** rest' with a bold coloured lead, as the reference does."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.line_spacing = 1.12
    if "**" in body:
        lead, rest = body.split("**", 2)[1], body.split("**", 2)[2]
        r1 = p.add_run(); r1.text = lead
        r1.font.size = Pt(size); r1.font.bold = True
        r1.font.name = D.FONT_BODY; r1.font.color.rgb = rgb(lead_colour or D.GOLD)
        r2 = p.add_run(); r2.text = rest
        r2.font.size = Pt(size); r2.font.name = D.FONT_BODY
        r2.font.color.rgb = rgb(colour)
    else:
        r = p.add_run(); r.text = body
        r.font.size = Pt(size); r.font.name = D.FONT_BODY
        r.font.color.rgb = rgb(colour)
    return tb


def text(s, x, y, w, h, body, size, colour, bold=False, mono=False,
         align=PP_ALIGN.LEFT, spacing=1.0, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(str(body).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = D.FONT_MONO if mono else D.FONT_BODY
        r.font.color.rgb = rgb(colour)
    return tb


def chrome(s):
    """Red hairline across the top — the reference deck's signature."""
    rect(s, 0, 0, W, 0.055, D.RED)


def header(s, kicker, title):
    """Returns the y (inches) where content may begin."""
    chrome(s)
    y = 0.42
    if kicker:
        text(s, M + 0.18, y, W - 2 * M, 0.28, kicker.upper(), 12, D.GOLD, bold=True)
        y += 0.30
    if title:
        rect(s, M, y + 0.04, 0.055, 0.42, D.GOLD)        # gold tick beside the title
        text(s, M + 0.18, y, W - 2 * M - 0.2, 0.62, title, 28, D.TEXT, bold=True, spacing=0.95)
        y += 0.60 * (1 + (len(title) > 52))
    return y + 0.30


def content_bottom():
    return H * D.CAPTION_TOP_FRAC


def phone(s, x, y, h, img):
    """Screenshot in a gold-outlined device frame, as in the reference deck."""
    from PIL import Image as _I
    iw, ih = _I.open(img).size
    w = h * iw / ih
    rect(s, x - 0.045, y - 0.045, w + 0.09, h + 0.09, D.GOLD)
    rect(s, x - 0.03, y - 0.03, w + 0.06, h + 0.06, D.BG)
    s.shapes.add_picture(img, Inches(x), Inches(y), Inches(w), Inches(h))
    return w


def L_stats(s, sp):
    """Big-number cards — the reference deck's 'By the numbers' grammar."""
    y = header(s, sp.get("kicker"), sp.get("title"))
    cards = sp["stats"][:6]
    per = 3 if len(cards) > 2 else len(cards)
    rows = [cards[i:i + per] for i in range(0, len(cards), per)]
    cw = (W - 2 * M - 0.28 * (per - 1)) / per
    ch = min(2.05, (content_bottom() - y - 0.28 * (len(rows) - 1)) / len(rows))
    for row in rows:
        for i, c in enumerate(row):
            x = M + i * (cw + 0.28)
            rect(s, x, y, cw, ch, D.PANEL)
            rect(s, x, y, cw, 0.05, D.RED)
            text(s, x + 0.28, y + 0.22, cw - 0.5, 0.9, str(c[0]), 46, D.GOLD, bold=True)
            text(s, x + 0.28, y + 1.05, cw - 0.5, 0.4, c[1], 15, D.TEXT, bold=True)
            if len(c) > 2 and c[2]:
                text(s, x + 0.28, y + 1.45, cw - 0.5, 0.4, c[2], 11, D.MUTED)
        y += ch + 0.28


def L_shot(s, sp):
    """Screenshot(s) beside supporting points."""
    y = header(s, sp.get("kicker"), sp.get("title"))
    imgs = sp["images"][:3]
    avail_h = content_bottom() - y
    if sp.get("items"):
        ph = avail_h - 0.15
        gap = 0.22
        total = 0
        for im in imgs:
            from PIL import Image as _I
            iw, ih = _I.open(im).size
            total += ph * iw / ih + gap
        x = W - M - total + gap
        for im in imgs:
            x += phone(s, x, y, ph, im) + gap
        yy = y + 0.1
        for b in sp["items"][:5]:
            rect(s, M, yy + 0.11, 0.11, 0.11, D.GOLD)
            text(s, M + 0.3, yy, W - 2 * M - total - 0.5, 0.8, b, 14, D.TEXT)
            yy += 0.44 * (1 + (len(b) > 52))
    else:
        ph = avail_h - 0.05
        widths = []
        from PIL import Image as _I
        for im in imgs:
            iw, ih = _I.open(im).size
            widths.append(ph * iw / ih)
        total = sum(widths) + 0.3 * (len(imgs) - 1)
        x = (W - total) / 2
        for im, w in zip(imgs, widths):
            phone(s, x, y, ph, im)
            x += w + 0.3
    if sp.get("note"):
        text(s, M, content_bottom() + 0.06, W - 2 * M, 0.4, sp["note"], 14, D.GOLD, bold=True)


# ------------------------------------------------------------------ layouts

def L_title(s, sp):
    chrome(s)
    rect(s, M, 2.30, 1.5, 0.05, D.GOLD)
    text(s, M, 1.35, W - 2 * M, 0.6, sp.get("kicker", "").upper(), 14, D.GOLD, bold=True)
    text(s, M, 2.60, W - 2 * M, 1.7, sp["title"], 44, D.TEXT, bold=True, spacing=0.95)
    if sp.get("sub"):
        text(s, M, 4.55, W - 2 * M, 0.9, sp["sub"], 17, D.MUTED)


def L_section(s, sp):
    chrome(s)
    rect(s, 0, 0.055, 2.55, H - 0.055, D.PANEL)
    text(s, 0.55, 2.55, 2.0, 1.6, sp["num"], 88, D.RED, bold=True)
    text(s, 0.55, 6.55, 2.0, 0.4, "TICKETPLEASE", 10, D.RULE, bold=True)
    rect(s, 3.15, 2.95, 0.055, 0.62, D.GOLD)
    text(s, 3.36, 2.95, W - 4.2, 1.0, sp["title"], 34, D.TEXT, bold=True)
    if sp.get("sub"):
        text(s, 3.36, 3.85, W - 4.2, 0.8, sp["sub"], 15, D.MUTED)


def L_statement(s, sp):
    chrome(s)
    text(s, M, 2.3, W - 2 * M, 2.2, sp["title"], 38, D.TEXT, bold=True, spacing=1.05)
    if sp.get("sub"):
        rect(s, M, 4.75, 1.2, 0.04, D.GOLD)
        text(s, M, 5.0, W - 2 * M, 0.9, sp["sub"], 16, D.MUTED)


def L_points(s, sp):
    y = header(s, sp.get("kicker"), sp.get("title"))
    items = sp["items"][:6]
    avail = content_bottom() - y
    tw = W - 2 * M
    if sp.get("image"):
        ph = avail - 0.1
        from PIL import Image as _I
        iw, ih = _I.open(sp["image"]).size
        pw = ph * iw / ih
        phone(s, W - M - pw, y + max(0, (avail - ph) / 2), ph, sp["image"])
        tw = W - 2 * M - pw - 0.5
    step = min(0.92, avail / max(1, len(items)))
    y += max(0, (avail - step * len(items)) / 2)
    for it in items:
        rect(s, M, y + 0.14, 0.12, 0.12, D.GOLD)
        rich(s, M + 0.34, y, tw - 0.34, step, it, 17, D.TEXT)
        y += step


def L_code(s, sp):
    y = header(s, sp.get("kicker"), sp.get("title"))
    lines = sp["code"].split("\n")
    fs = sp.get("pptsize", 13)
    bh = len(lines) * fs * 1.62 / 72 + 0.36
    avail = content_bottom() - y
    y += max(0, (avail - bh - (0.5 if sp.get("note") else 0)) / 2)
    if sp.get("file"):
        text(s, M, y - 0.34, W - 2 * M, 0.3, sp["file"], 11, D.MUTED, mono=True)
    rect(s, M, y, W - 2 * M, bh, D.PANEL)
    rect(s, M, y, 0.05, bh, D.GOLD)
    text(s, M + 0.28, y + 0.18, W - 2 * M - 0.5, bh, sp["code"], fs, D.TEXT,
         mono=True, spacing=1.35)
    if sp.get("note"):
        text(s, M, y + bh + 0.24, W - 2 * M, 0.5, sp["note"], 16, D.GOLD, bold=True)


def L_compare(s, sp):
    y = header(s, sp.get("kicker"), sp.get("title"))
    cols = sp["cols"][:2]
    cw = (W - 2 * M - 0.36) / 2

    def ch_of(c):
        h = 0.9
        if c.get("code"):
            h += len(c["code"].split("\n")) * 0.245 + 0.2
        for b in c.get("items", [])[:4]:
            h += 0.42 * (1 + (len(b) > 40))
        return h + 0.24

    bh = max(2.6, max(ch_of(c) for c in cols))
    avail = content_bottom() - y - (0.62 if sp.get("verdict") else 0)
    y += max(0, (avail - bh) / 2)
    for i, c in enumerate(cols):
        x = M + i * (cw + 0.36)
        accent = D.RED if c.get("warn") else (D.GREEN if c.get("ok") else D.GOLD)
        rect(s, x, y, cw, bh, D.PANEL)
        rect(s, x, y, 0.05, bh, accent)
        text(s, x + 0.3, y + 0.26, cw - 0.5, 0.5, c["head"], 21, accent, bold=True)
        yy = y + 0.86
        if c.get("code"):
            n = len(c["code"].split("\n"))
            text(s, x + 0.3, yy, cw - 0.5, n * 0.24, c["code"], 11, D.TEXT,
                 mono=True, spacing=1.35)
            yy += n * 0.245 + 0.2
        for b in c.get("items", [])[:4]:
            text(s, x + 0.3, yy, cw - 0.5, 0.6, b, 15, D.MUTED)
            yy += 0.42 * (1 + (len(b) > 40))
    if sp.get("verdict"):
        text(s, M, y + bh + 0.2, W - 2 * M, 0.5, sp["verdict"], 17, D.GOLD, bold=True)


def L_chips(s, sp):
    y = header(s, sp.get("kicker"), sp.get("title"))
    chips = [(c if isinstance(c, (list, tuple)) else (c, "")) for c in sp["chips"][:8]]
    per = 4 if len(chips) > 3 else max(1, len(chips))
    rows = [chips[i:i + per] for i in range(0, len(chips), per)]
    cw = (W - 2 * M - 0.3 * (per - 1)) / per
    ch = 1.42
    avail = content_bottom() - y
    y += max(0, (avail - len(rows) * (ch + 0.24)) / 2)
    for row in rows:
        for i, (term, mean) in enumerate(row):
            x = M + i * (cw + 0.3)
            rect(s, x, y, cw, ch, D.PANEL)
            rect(s, x, y, cw, 0.05, D.GOLD)
            # shrink long terms so they can't spill out of the tile
            avail_in = cw - 0.48
            size = 19
            while size > 11 and len(term) * size * 0.60 / 72 > avail_in:
                size -= 1
            text(s, x + 0.24, y + 0.26, cw - 0.4, 0.42, term, size, D.TEXT, bold=True, mono=True)
            if mean:
                text(s, x + 0.24, y + 0.74, cw - 0.4, 0.36, mean, 12, D.MUTED)
        y += ch + 0.24


def L_steps(s, sp):
    y = header(s, sp.get("kicker"), sp.get("title"))
    items = sp["items"][:5]
    rh = 0.94
    avail = content_bottom() - y
    y += max(0, (avail - len(items) * (rh + 0.16)) / 2)
    for i, it in enumerate(items, 1):
        rect(s, M, y, W - 2 * M, rh, D.PANEL)
        rect(s, M, y, 0.05, rh, D.GOLD)
        text(s, M + 0.3, y + 0.22, 0.7, 0.4, f"{i:02d}", 16, D.GOLD, bold=True, mono=True)
        text(s, M + 1.05, y + 0.2, W - 2 * M - 1.4, 0.5, it, 16, D.TEXT)
        y += rh + 0.16


def L_table(s, sp):
    y = header(s, sp.get("kicker"), sp.get("title"))
    rows = sp["rows"][:7]
    rh = min(0.78, (content_bottom() - y) / max(1, len(rows)))
    y += max(0, (content_bottom() - y - rh * len(rows)) / 2)
    # size the key column to the widest key, shrinking the font before overflowing
    longest = max(len(str(a)) for a, _ in rows)
    ksize = 14
    kw = longest * ksize * 0.60 / 72
    while kw > 4.4 and ksize > 9:
        ksize -= 1
        kw = longest * ksize * 0.60 / 72
    kw = max(1.0, kw + 0.12)
    for i, (k, v) in enumerate(rows):
        rect(s, M, y, W - 2 * M, rh - 0.04, D.PANEL if i % 2 == 0 else D.PANEL_2)
        rect(s, M, y, 0.05, rh - 0.04, D.GOLD)
        text(s, M + 0.3, y + rh / 2 - 0.15, kw, 0.4, k, ksize, D.GOLD, bold=True, mono=True)
        text(s, M + 0.3 + kw + 0.3, y + rh / 2 - 0.15, W - 2 * M - kw - 0.9, 0.4, v, 14, D.TEXT)
        y += rh


LAYOUTS = dict(title=L_title, section=L_section, statement=L_statement, points=L_points,
               code=L_code, compare=L_compare, chips=L_chips, steps=L_steps, table=L_table,
               stats=L_stats, shot=L_shot)


def build(module, specs, out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    made = 0
    for sp in specs:
        if sp is None:
            continue
        LAYOUTS[sp["layout"]](blank(prs), sp)
        made += 1
    prs.save(out_path)
    return made
